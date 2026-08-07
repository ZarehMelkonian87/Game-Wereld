import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
import os

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="D1DCD6"):
    tblPr = table._tbl.tblPr
    borders_elm = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{color}"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders_elm)

def add_callout_box(doc, text_list, title="UX DESIGN NOTE", border_color="0B8457", bg_color="F4F7F6"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(6.5)
    cell = tbl.cell(0, 0)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=140)
    
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:top w:val="none"/>'
        f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
        f'  <w:bottom w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    run_t = p.add_run(f"📌 {title}\n")
    run_t.bold = True
    run_t.font.name = "Calibri"
    run_t.font.size = Pt(10.5)
    run_t.font.color.rgb = RGBColor(11, 132, 87) if border_color=="0B8457" else RGBColor(217, 131, 16)
    
    for item in text_list:
        p_item = cell.add_paragraph()
        p_item.paragraph_format.space_after = Pt(3)
        p_item.paragraph_format.line_spacing = 1.15
        run_i = p_item.add_run(item)
        run_i.font.name = "Calibri"
        run_i.font.size = Pt(9.5)
        run_i.font.color.rgb = RGBColor(40, 50, 60)
        
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def add_screenshot_figure(doc, img_path, caption_text):
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=Inches(2.6))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(10)
        run_cap = p_cap.add_run(f"Figuur: {caption_text}")
        run_cap.font.name = "Calibri"
        run_cap.font.size = Pt(9)
        run_cap.font.italic = True
        run_cap.font.color.rgb = RGBColor(90, 106, 117)

def create_table(doc, headers, data):
    t = doc.add_table(rows=len(data)+1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    
    col_widths = [Inches(1.6), Inches(1.3), Inches(1.7), Inches(1.0), Inches(0.9)]
    for i, w in enumerate(col_widths):
        if i < len(headers):
            t.columns[i].width = w
            
    for i, h_text in enumerate(headers):
        c = t.rows[0].cells[i]
        set_cell_background(c, "0B8457")
        set_cell_margins(c, top=80, bottom=80, left=80, right=80)
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h_text)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9)

    TEAL = RGBColor(11, 132, 87)
    for r_idx, row in enumerate(data, start=1):
        cells = t.rows[r_idx].cells
        bg = "F4F7F6" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row):
            set_cell_background(cells[c_idx], bg)
            set_cell_margins(cells[c_idx], top=60, bottom=60, left=80, right=80)
            p = cells[c_idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(val)
            r.font.size = Pt(8.5)
            if c_idx == 0:
                r.font.bold = True
                r.font.color.rgb = TEAL

    set_table_borders(t)
    doc.add_paragraph().paragraph_format.space_after = Pt(14)

def build_docx():
    doc = Document()
    
    brain_dir = '/Users/melkonian/.gemini/antigravity-ide/brain/c6315284-cea7-4995-a87b-752c0862c7bd'
    
    img_main = os.path.join(brain_dir, 'media__1784836812683.png')
    img_select = os.path.join(brain_dir, 'media__1784836812605.png')
    img_settings = os.path.join(brain_dir, 'media__1784836812579.png')
    img_reward = os.path.join(brain_dir, 'media__1784836812572.png')
    img_fly_start = os.path.join(brain_dir, 'media__1784879932153.png')
    img_kies_woord = os.path.join(brain_dir, 'media__1784879932176.png')
    img_zeg_zet_kbd = os.path.join(brain_dir, 'media__1784879932198.png')
    img_zeg_zet_game = os.path.join(brain_dir, 'media__1784879932221.png')
    img_fly_active = os.path.join(brain_dir, 'media__1784879932236.png')

    # Margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(26, 37, 44)
    
    TEAL = RGBColor(11, 132, 87)
    GOLD = RGBColor(217, 131, 16)
    GRAY = RGBColor(90, 106, 117)
    
    # Title Banner
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    r_sub = p_title.add_run("COMPLETE GAME UX/UI SPECIFICATION DOCUMENT (ALL 9 SCREENS)")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(11)
    r_sub.font.bold = True
    r_sub.font.color.rgb = GOLD
    
    p_main = doc.add_paragraph()
    p_main.paragraph_format.space_after = Pt(6)
    r_m = p_main.add_run("Strand-bezem-escape\n\"Magisch Strand Avontuur\"")
    r_m.font.name = "Calibri"
    r_m.font.size = Pt(24)
    r_m.font.bold = True
    r_m.font.color.rgb = TEAL
    
    p_desc = doc.add_paragraph()
    p_desc.paragraph_format.space_after = Pt(14)
    r_d = p_desc.add_run("Volledige Interface Componenten, Naming Conventions, Gameplay Flows & Modal Specificatie (9 Schermen)")
    r_d.font.name = "Calibri"
    r_d.font.size = Pt(12)
    r_d.font.italic = True
    r_d.font.color.rgb = GRAY

    # Metadata Table
    meta_table = doc.add_table(rows=5, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False
    meta_table.columns[0].width = Inches(2.0)
    meta_table.columns[1].width = Inches(4.5)
    
    meta_data = [
        ("Projectnaam:", "Strand-bezem-escape (Magisch Strand Avontuur)"),
        ("Document Versie:", "v2.0 (Volledige Schermspecificatie - 9 Schermen)"),
        ("Rol / Auteur:", "Senior Game UX/UI Designer & Technical Product Owner"),
        ("Aantal Schermen:", "9 Unieke Schermen & Modals (Menu's, Modals, Gameplay, Overlays)"),
        ("Doelgroep:", "Game Developers, UI/UX Designers, Edu-Tech Content Creators, QA Testers")
    ]
    
    for idx, (label, val) in enumerate(meta_data):
        cell_lbl = meta_table.cell(idx, 0)
        cell_val = meta_table.cell(idx, 1)
        set_cell_background(cell_lbl, "F4F7F6")
        set_cell_background(cell_val, "FFFFFF")
        set_cell_margins(cell_lbl, top=60, bottom=60, left=100, right=100)
        set_cell_margins(cell_val, top=60, bottom=60, left=100, right=100)
        
        p_l = cell_lbl.paragraphs[0]
        p_l.paragraph_format.space_after = Pt(0)
        r_l = p_l.add_run(label)
        r_l.bold = True
        r_l.font.size = Pt(9.5)
        
        p_v = cell_val.paragraphs[0]
        p_v.paragraph_format.space_after = Pt(0)
        r_v = p_v.add_run(val)
        r_v.font.size = Pt(9.5)

    set_table_borders(meta_table)
    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    
    # ---------------------------------------------------------
    # SECTION 1: INLEIDING & DOELSTELLING
    # ---------------------------------------------------------
    h1 = doc.add_heading(level=1)
    p_h1 = h1.paragraphs[0] if hasattr(h1, 'paragraphs') else h1
    r = p_h1.add_run("1. Inleiding & UX Ontwerpprincipes")
    r.font.name = "Calibri"
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = TEAL
    p_h1.paragraph_format.space_before = Pt(12)
    p_h1.paragraph_format.space_after = Pt(6)
    
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    p.add_run(
        "Dit document bevat het volledige en geharmoniseerde User Experience (UX) en User Interface (UI) "
        "ontwerpdocument voor alle 9 schermen en overlays van Strand-bezem-escape (Magisch Strand Avontuur). "
        "Met alle menu's, gameplay-omgevingen, instructiemodals, invoeroverlays en beloningsschermen vastgelegd "
        "in gestandaardiseerde tabellen, vormt dit de definitieve blauwdruk voor ontwikkeling en test-acceptatie."
    )
    
    add_callout_box(
        doc,
        [
            "• Kindvriendelijke Ergonomie: Knoppen hebben een minimale touch-target van 48x48dp met duidelijke iconografie.",
            "• Directe Multimodale Feedback: Elk interactief element reageert visueel (schaalverandering/pulse) en auditief (klank/click).",
            "• Multimodale Invoer: Ondersteuning voor gesproken antwoorden (microfoon), slepen & neerzetten (drag & drop) en tekstinvoer.",
            "• Inclusiviteit & Privacy: Geen bewaarde spraakopnames, opties voor rustige beweging (reduced motion) en directe toegankelijkheid."
        ],
        title="KERNPRINCIPES VAN STRAND-BEZEM-ESCAPE UX"
    )

    # ---------------------------------------------------------
    # SECTION 2: UI NAAMGEVING EN DESIGN SYSTEEM STANDAARD
    # ---------------------------------------------------------
    h1 = doc.add_heading(level=1)
    p_h1 = h1.paragraphs[0] if hasattr(h1, 'paragraphs') else h1
    r = p_h1.add_run("2. Gestandaardiseerde UI Naming Conventions")
    r.font.name = "Calibri"
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = TEAL
    p_h1.paragraph_format.space_before = Pt(14)
    p_h1.paragraph_format.space_after = Pt(6)

    prefix_table = doc.add_table(rows=11, cols=3)
    prefix_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    prefix_table.autofit = False
    prefix_table.columns[0].width = Inches(1.5)
    prefix_table.columns[1].width = Inches(1.8)
    prefix_table.columns[2].width = Inches(3.2)
    
    headers = ["Prefix", "Component Type", "Voorbeeld ID"]
    hdr_cells = prefix_table.rows[0].cells
    for i, h_text in enumerate(headers):
        set_cell_background(hdr_cells[i], "0B8457")
        set_cell_margins(hdr_cells[i], top=80, bottom=80, left=100, right=100)
        p = hdr_cells[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h_text)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9.5)
        
    prefix_data = [
        ("SCR_", "Scherm / View", "SCR_MAIN_TITLE, SCR_ZEG_ZET_GAME"),
        ("BTN_", "Interactieve Knop", "BTN_PRIMARY_PLAY, BTN_ACTION_KLAAR"),
        ("CARD_", "Selectie-, Modal- of Informatiekaart", "CARD_GAME_ZEG_ZET, CARD_MODAL_KEYBOARD"),
        ("TOGGLE_", "Aan/Uit Schakelaar", "TOGGLE_AUDIO, TOGGLE_REDUCED_MOTION"),
        ("DSP_", "Weergave / Statusteller", "DSP_STAR_COUNTER, DSP_DISTANCE_COUNTER"),
        ("TTL_", "Titel / Header Capsule", "TTL_HEADER_PILL, TTL_MODAL_TITLE"),
        ("INFOBOX_", "Melding / Waarschuwingsvak", "INFOBOX_SPEECH_STATUS, INFOBOX_VIDEO_ERROR"),
        ("CANVAS_", "Interactief Spelveld / Substraat", "CANVAS_BEACH_OCEAN, CANVAS_FLIGHT_ARENA"),
        ("TRAY_", "Asset Palet / Keuzebalk", "TRAY_STICKER_PALETTE"),
        ("DOCK_", "Onderste Actie- / Navigatiebalk", "DOCK_SELECTION_FOOTER, DOCK_FLIGHT_CONTROLS")
    ]
    
    for row_idx, data in enumerate(prefix_data, start=1):
        row_cells = prefix_table.rows[row_idx].cells
        bg = "F4F7F6" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(data):
            set_cell_background(row_cells[col_idx], bg)
            set_cell_margins(row_cells[col_idx], top=60, bottom=60, left=100, right=100)
            p = row_cells[col_idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            r.font.size = Pt(9)
            if col_idx == 0:
                r.font.bold = True
                r.font.color.rgb = TEAL

    set_table_borders(prefix_table)
    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    std_headers = ["Element Code Name", "Visuele Naam", "Functionaliteit & Omschrijving", "Visuele Stijl", "Doel / Actie"]

    # ---------------------------------------------------------
    # SECTION 3: SCHERM 1 - HOOFDSCHERM (SCR_MAIN_TITLE)
    # ---------------------------------------------------------
    h1 = doc.add_heading(level=1)
    p_h1 = h1.paragraphs[0] if hasattr(h1, 'paragraphs') else h1
    r = p_h1.add_run("3. Scherm 1: Hoofdscherm / Titelmenu (SCR_MAIN_TITLE)")
    r.font.name = "Calibri"
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = TEAL
    p_h1.paragraph_format.space_before = Pt(14)
    p_h1.paragraph_format.space_after = Pt(6)

    add_screenshot_figure(doc, img_main, "SCR_MAIN_TITLE - Hoofdscherm / Titelmenu Interface")
    s1_data = [
        ("BTN_NAV_BACK", "Terug Knop", "Sluit de game en keert terug naar de overkoepelende app-omgeving.", "Witte cirkel knop met donkere pijl-links", "Navigate Back"),
        ("BTN_AUDIO_TOGGLE_QUICK", "Snelle Audio Knop", "Schakelt audio/geluid direct in of uit vanuit het hoofdscherm.", "Witte cirkel knop met luidspreker-icoon", "Toggle Audio"),
        ("DSP_STAR_COUNTER", "Sterrenteller", "Toont het huidige aantal verzamelde beloningssterren van de speler (bijv. 120).", "Witte capsule pil met gouden rand en ster", "Read Only Display"),
        ("BTN_SETTINGS_GEAR", "Instellingen Knop", "Opent het uitgebreide instellingen- en privacy-menu.", "Gele afgeronde knop met donker tandwiel", "Open SCR_SETTINGS_PRIVACY"),
        ("IMG_GAME_LOGO", "Titel Logo", "Visueel logo 'MAGISCH STRAND AVONTUUR' met 3D-effecten en strandthema.", "Gele/Blauwe 3D typografie met decoratie", "Visual Branding"),
        ("IMG_HERO_CHARACTER", "Mascotte Illustratie", "Hero-illustratie van de jongen op de vliegende strandbezem met de regenboogster.", "Kleurrijke karakter-art op strandachtergrond", "Visual Engagement"),
        ("BTN_PRIMARY_PLAY", "Spelen Knop", "Hoofdactieknop om het spel te starten en het avontuur te kiezen.", "Grote groen-gebolde pilknop met Play-icoon", "Open SCR_ADVENTURE_SELECT")
    ]
    create_table(doc, std_headers, s1_data)

    # ---------------------------------------------------------
    # SECTION 4: SCHERM 2 - SPEL SELECTIESCHERM (SCR_ADVENTURE_SELECT)
    # ---------------------------------------------------------
    h1 = doc.add_heading(level=1)
    p_h1 = h1.paragraphs[0] if hasattr(h1, 'paragraphs') else h1
    r = p_h1.add_run("4. Scherm 2: Spel Selectiescherm / Kies Avontuur (SCR_ADVENTURE_SELECT)")
    r.font.name = "Calibri"
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = TEAL
    p_h1.paragraph_format.space_before = Pt(14)
    p_h1.paragraph_format.space_after = Pt(6)

    add_screenshot_figure(doc, img_select, "SCR_ADVENTURE_SELECT - Spel Selectiescherm Interface")
    s2_data = [
        ("BTN_NAV_BACK", "Terug Knop", "Navigeert terug naar het Hoofdscherm.", "Witte cirkel knop met pijl-links", "Open SCR_MAIN_TITLE"),
        ("TTL_HEADER_PILL", "Header Titel", "Schermtitel capsule met de tekst 'Kies avontuur'.", "Witte afgeronde pill-header", "Static Title"),
        ("DSP_STAR_COUNTER", "Sterrenteller", "Toont actueel aantal verdiende sterren (120).", "Witte capsule pil met ster", "Read Only Display"),
        ("LBL_SECTION_TITLE", "Sectielabel", "Subtitel boven de spellijst: 'Kies spel'.", "Donkergrijze vetgedrukte tekst", "Section Divider"),
        ("CARD_GAME_ZEG_ZET", "Spelkaart: Zeg & Zet", "Selecteert de luister-, spreek- of typ-opdracht minigame.", "Witte afgeronde kaart met groene rand & spreekwolk-icoon", "Select Minigame 1"),
        ("CARD_GAME_KIES_WOORD", "Spelkaart: Kies het Woord", "Selecteert de meervoudige keuze minigame.", "Witte kaart met lichtblauwe rand & boek-icoon", "Select Minigame 2"),
        ("CARD_GAME_ZEG_VLIEG", "Spelkaart: Zeg & Vlieg", "Selecteert de spraakgestuurde bezem-vlieggame.", "Zachtgele kaart met gouden rand & microfoon-icoon", "Select Minigame 3"),
        ("BTN_PRIMARY_START", "Start Spel Knop", "Start direct de op dat moment geselecteerde minigame.", "Brede groene knop onderaan met Play-icoon", "Launch Active Game"),
        ("BTN_SECONDARY_REWARD", "Beloning Knop", "Navigeert naar het resultatenoverzicht en het stickerboek.", "Creme knop met gele rand en cadeau-icoon", "Open SCR_REWARD_SUMMARY"),
        ("BTN_SECONDARY_OPTIONS", "Opties Knop", "Opent het instellingenmenu voor geluid, hints en devtools.", "Witte knop met blauwe/grijze rand en tandwiel", "Open SCR_SETTINGS_PRIVACY")
    ]
    create_table(doc, std_headers, s2_data)

    # ---------------------------------------------------------
    # SECTION 5: SCHERM 3 - INSTELLINGEN & PRIVACY (SCR_SETTINGS_PRIVACY)
    # ---------------------------------------------------------
    h1 = doc.add_heading(level=1)
    p_h1 = h1.paragraphs[0] if hasattr(h1, 'paragraphs') else h1
    r = p_h1.add_run("5. Scherm 3: Instellingen & Privacy Menu (SCR_SETTINGS_PRIVACY)")
    r.font.name = "Calibri"
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = TEAL
    p_h1.paragraph_format.space_before = Pt(14)
    p_h1.paragraph_format.space_after = Pt(6)

    add_screenshot_figure(doc, img_settings, "SCR_SETTINGS_PRIVACY - Instellingen & Privacy Interface")
    s3_data = [
        ("BTN_NAV_MENU", "Menu Knop", "Sluit het instellingenscherm en keert terug naar de vorige pagina.", "Witte pilknop met pijl-links en tekst 'Menu'", "Return to Sender"),
        ("TTL_HEADER_PILL", "Header Titel", "Schermtitel capsule met de tekst 'Instellingen'.", "Witte afgeronde capsule header", "Static Title"),
        ("TOGGLE_AUDIO", "Audio Schakelaar", "Schakelt gesproken opdrachten en video-audio in/uit.", "Witte kaart met luidspreker-icoon en groene Toggle Switch", "Toggle Audio State"),
        ("TOGGLE_MUSIC", "Muziek Schakelaar", "Schakelt de achtergrondmuziek in/uit.", "Witte kaart met muzieknoot-icoon en grijze Toggle Switch", "Toggle BGM State"),
        ("TOGGLE_HINTS", "Hints Schakelaar", "Schakelt mascottesubsidie en automatische hints in/uit.", "Witte kaart met lamp-icoon en groene Toggle Switch", "Toggle Hint System"),
        ("TOGGLE_REDUCED_MOTION", "Rustige Beweging Toggle", "Vermindert animaties en pulse-effecten voor rustige ervaring.", "Witte kaart met oog-kruis icoon en grijze Toggle Switch", "Toggle Accessibility"),
        ("TOGGLE_DEVTOOLS", "Zone Editor Toggle", "Opent interactieve zone-locatie editor voor ontwikkelaars.", "Witte kaart met sleutel-icoon en grijze Toggle Switch", "Developer Mode"),
        ("CARD_PRIVACY_SECTION", "Privacy Container", "Gele achtergrondkaart die alle microfooninformatie bundelt.", "Gele afgeronde container met schild-icoon", "Grouping Container"),
        ("LBL_PRIVACY_DESC", "Privacy Uitleg", "Omschrijving: 'De microfoon wordt alleen gebruikt om korte zinnen naar tekst om te zetten.'", "Donkergrijze tekst onder privacy-titel", "Informational Text"),
        ("DROPDOWN_PRIVACY_FAQ", "FAQ Dropdown", "Uitklapbare knop: 'Waarom gebruiken we de microfoon?'", "Lichtblauwe pilknop met dropdown pijl", "Toggle FAQ Details"),
        ("INFOBOX_SPEECH_STATUS", "Status Infoblok", "Blauw vak met status en alternatief (typen bij storing).", "Lichtblauw afgerond vak met telefoon-icoon", "Status Notification"),
        ("INFOBOX_PERMISSION_NOTICE", "Toestemming Infoblok", "Geel vak over browser-popup toestemming.", "Lichtgeel vak met oranje rand", "Permission Alert"),
        ("BTN_MIC_RECHECK", "Controleer Opnieuw Knop", "Herstart de spraakherkenningstest en vraagt opnieuw toestemming.", "Groene brede knop met microfoon-icoon", "Trigger Mic Permission"),
        ("LBL_MIC_INSTRUCTION", "Instructietekst", "Ondersteunende tekst: 'Tik op de knop om microfoontoegang te vragen.'", "Kleine grijze instructietekst", "Instructional Label"),
        ("INFOBOX_MIC_BLOCKED_ALERT", "Geblokkeerd Waarschuwing", "Waarschuwingsvak bij geblokkeerde microfoon.", "Oranje/Rood omrand waarschuwingsvak", "Error State Alert")
    ]
    create_table(doc, std_headers, s3_data)

    # ---------------------------------------------------------
    # SECTION 6: SCHERM 4 - BELONING & RESULTATEN (SCR_REWARD_SUMMARY)
    # ---------------------------------------------------------
    h1 = doc.add_heading(level=1)
    p_h1 = h1.paragraphs[0] if hasattr(h1, 'paragraphs') else h1
    r = p_h1.add_run("6. Scherm 4: Beloning & Resultaten Scherm (SCR_REWARD_SUMMARY)")
    r.font.name = "Calibri"
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = TEAL
    p_h1.paragraph_format.space_before = Pt(14)
    p_h1.paragraph_format.space_after = Pt(6)

    add_screenshot_figure(doc, img_reward, "SCR_REWARD_SUMMARY - Beloning & Resultaten Interface")
    s4_data = [
        ("TTL_HEADER_PILL", "Header Titel", "Schermtitel capsule met de tekst 'Beloning'.", "Witte afgeronde capsule header", "Static Title"),
        ("CARD_REWARD_SHOWCASE", "Sticker Showcase Kaart", "Centraal kader dat de verdiende sticker (Schelp Sticker) en regenboog-ster badge toont.", "Creme achtergrond met gele rand en sticker-art", "Reward Showcase"),
        ("DSP_STICKER_PROGRESS_PILL", "Sticker Voortgangsbadge", "Toont de voortgang van de stickerverzameling ('0/30' met ster).", "Gele afgeronde pil met ster-icoon", "Progress Indicator"),
        ("LBL_STICKER_NAME", "Sticker Naam Label", "Tekstlabel onder de stickerweergave: 'Schelp Sticker'.", "Vetgedrukte donkere tekst", "Item Identification"),
        ("STAT_BOX_GOED", "Statistiek: GOED", "Toont het aantal foutloos beantwoorde vragen ('GOED: 0').", "Lichtgroene capsule pil met groene rand", "Score Counter"),
        ("STAT_BOX_TEMPO", "Statistiek: TEMPO", "Toont de behaalde snelheidsbonuspuntenscore ('TEMPO: +0').", "Lichtblauwe capsule pil met blauwe rand", "Speed Bonus Counter"),
        ("STAT_BOX_HINTS", "Statistiek: HINTS", "Toont het aantal geraadpleegde hints tijdens de sessie ('HINTS: 0').", "Lichtoranje capsule pil met oranje rand", "Assist Counter"),
        ("STAT_BOX_AUDIO", "Statistiek: AUDIO", "Toont het aantal keren dat geluid/opdracht is herhaald ('AUDIO: 0').", "Witte capsule pil met grijze rand", "Audio Replay Counter"),
        ("STAT_BOX_STERREN", "Statistiek: STERREN", "Toont het totale aantal netto gewonnen beloningssterren ('STERREN: +0').", "Zachtgouden capsule pil met gouden rand", "Total Currency Awarded"),
        ("SEC_WORDS_PRACTICED", "Woorden Sectie", "Overzicht van geoefende woorden. Status: 'nog geen woorden'.", "Titel met pill-tag 'nog geen woorden'", "Vocabulary Summary"),
        ("SEC_SPATIAL_WORDS", "Plaatswoorden Sectie", "Overzicht van geoefende ruimtelijke plaatswoorden. Status: 'nog geen plaatswoorden'.", "Titel met pill-tag 'nog geen plaatswoorden'", "Grammar/Spatial Summary"),
        ("BANNER_REWARD_SUMMARY", "Beloning Samenvatting", "Onderste trofee-banner: 'Beloning: Schelp Sticker' met trofee-icoon.", "Gele afgeronde banner met trofee-icoon", "Reward Highlight"),
        ("BTN_ACTION_REPLAY", "Opnieuw Knop", "Herstart direct de zojuist gespeelde minigame of sessie.", "Groene capsule knop met refresh-icoon", "Replay Session"),
        ("BTN_ACTION_WORLD", "Wereld Knop", "Navigeert naar de overkoepelende Wereldkaart van Game-Wereld.", "Groene capsule knop met wereldbol-icoon", "Open World Map"),
        ("BTN_ACTION_MENU", "Menu Knop", "Navigeert terug naar het Hoofdscherm (SCR_MAIN_TITLE).", "Groene capsule knop met home-icoon", "Open Main Menu")
    ]
    create_table(doc, std_headers, s4_data)

    # ---------------------------------------------------------
    # SECTION 7: SCHERM 5 - ZEG & VLIEG START MODAL (SCR_ZEG_VLIEG_START)
    # ---------------------------------------------------------
    h1 = doc.add_heading(level=1)
    p_h1 = h1.paragraphs[0] if hasattr(h1, 'paragraphs') else h1
    r = p_h1.add_run("7. Scherm 5: Zeg & Vlieg Start & Instructie Modal (SCR_ZEG_VLIEG_START)")
    r.font.name = "Calibri"
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = TEAL
    p_h1.paragraph_format.space_before = Pt(14)
    p_h1.paragraph_format.space_after = Pt(6)

    add_screenshot_figure(doc, img_fly_start, "SCR_ZEG_VLIEG_START - Zeg & Vlieg Instructie & Start Modal")
    s5_data = [
        ("BTN_NAV_HOME", "Home Knop", "Verlaat de minigame en keert terug naar het avontuur-selectiescherm.", "Afgeronde vierkante knop met home-icoon", "Exit to Menu"),
        ("DSP_DISTANCE_COUNTER", "Afstandsbar", "Toont de afgelegde afstand en het doel (bijv. 'Afstand 0m', progress bar 0/100).", "Blauwe capsule met voortgangsbalk", "Distance Tracker"),
        ("DSP_LEVEL_TROPHY_BADGE", "Level & Score Badge", "Toont trofeebadge '0', sterren '0' en actueel level 'Level 1'.", "Blauwe pil met trofee en sterren-icoon", "Score & Level Display"),
        ("CARD_MODAL_ZEG_VLIEG_START", "Instructie Card Modal", "Centraal wit venster met de instructies en instellingen van Zeg & Vlieg.", "Witte kaart met afgeronde hoeken en schaduw", "Instruction Container"),
        ("BADGE_MODAL_STAR_HEADER", "Ster Mascotte Header", "Badge met vrolijke regenboogster-mascotte boven de titel.", "Gele afgeronde vierkante badge", "Visual Mascot Header"),
        ("TTL_MODAL_TITLE", "Modal Titel", "Titel van de minigame: 'Zeg & Vlieg'.", "Donkerblauwe vetgedrukte titel", "Minigame Title"),
        ("LBL_MODAL_INSTRUCTION", "Instructietekst", "Spelregels: 'Vlieg zo ver mogelijk. Noem plaatjes die je ziet. Raak geen obstakel.'", "Donkergrijze instructietekst", "Game Objective"),
        ("CONTAINER_TARGET_WORDS", "Doelwoorden Container", "Verzameling van te noemen strandwoorden met ster-iconen.", "Raster van gele pil-tags met sterren", "Target Vocabulary List"),
        ("CHIP_WORD_TAGS", "Doelwoord Chips", "Individuele chips: 'parasol', 'zon', 'schelp', 'dolfijn', 'krab', 'boot', 'bal'.", "Gele capsules met ster-icoon en tekst", "Vocabulary Tag Item"),
        ("INFOBOX_PRIVACY_NOTE", "Privacy Garantiestempel", "Groene pil met schild-icoon: 'We slaan geen opname op.'", "Groen afgerond vak met schild-icoon", "Privacy Assurance"),
        ("BTN_PRIMARY_START_FLY", "Start Knop", "Start direct de actieve vlieg-gameplay (SCR_ZEG_VLIEG_ACTIVE).", "Grote brede groene knop met Play-icoon", "Launch Flight Gameplay"),
        ("PANEL_FLY_CONTROLS_OVERLAY", "Onderste Besturingsbalk", "Informatiekaart 'Zeg & Vlieg' met actieknoppen voor handmatige besturing.", "Witte en gele containers onderaan", "Flight Controls Footer"),
        ("BTN_MANUAL_FLY_UP", "Omhoog Knop", "Beweegt de strandbezem handmatig omhoog als spraak niet wordt gebruikt.", "Groene knop met pijl-omhoog", "Manual Flight UP"),
        ("BTN_MANUAL_FLY_DOWN", "Omlaag Knop", "Beweegt de strandbezem handmatig omlaag.", "Groene knop met pijl-omlaag", "Manual Flight DOWN")
    ]
    create_table(doc, std_headers, s5_data)

    # ---------------------------------------------------------
    # SECTION 8: SCHERM 6 - KIES HET WOORD GAMEPLAY (SCR_KIES_WOORD_GAME)
    # ---------------------------------------------------------
    h1 = doc.add_heading(level=1)
    p_h1 = h1.paragraphs[0] if hasattr(h1, 'paragraphs') else h1
    r = p_h1.add_run("8. Scherm 6: Kies het Woord Quiz Gameplay (SCR_KIES_WOORD_GAME)")
    r.font.name = "Calibri"
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = TEAL
    p_h1.paragraph_format.space_before = Pt(14)
    p_h1.paragraph_format.space_after = Pt(6)

    add_screenshot_figure(doc, img_kies_woord, "SCR_KIES_WOORD_GAME - Kies het Woord Quiz Interface")
    s6_data = [
        ("BTN_NAV_BACK", "Terug Knop", "Verlaat de quiz en keert terug naar het spel-selectiescherm.", "Witte cirkel knop met pijl-links", "Navigate Back"),
        ("BTN_AUDIO_TOGGLE", "Audio Knop", "Schakelt gesproken vraag en geluidseffecten in of uit.", "Blauwe cirkel knop met luidspreker-icoon", "Toggle Audio"),
        ("DSP_STAR_COUNTER", "Sterrenteller", "Toont actueel verdiende beloningssterren ('0').", "Gele capsule pil met ster", "Score Display"),
        ("BTN_HINT_ASSIST", "Hint Knop", "Vraagt een visuele/gesproken hint aan bij de mascotte.", "Gele afgeronde knop met gloeilamp-icoon", "Trigger Hint"),
        ("CARD_QUESTION_PROMPT", "Vraagstelling Card", "Bovenvak met mascottesticker, de vraagtekst en herhaalknop.", "Witte afgeronde container met schaduw", "Question Banner"),
        ("IMG_MASCOT_SPEAKER", "Mascotte Badge", "Illustratie van de sterrenfeemascotte die de vraag uitspreekt.", "Ronde badge met regenboogster", "Visual Speaker Badge"),
        ("LBL_QUESTION_TEXT", "Vraagtekst", "Gesproken/geschreven vraag: 'Waar is de krab?'", "Donkerblauwe vetgedrukte vraagtekst", "Question Prompt"),
        ("BTN_AUDIO_REPLAY_PROMPT", "Herhaal Audio Knop", "Herhaalt het afspelen van de gesproken vraagzin.", "Blauwe cirkel knop met luidspreker-icoon", "Replay Question SFX"),
        ("GRID_CHOICE_CARDS", "Keuzekaarten Raster", "Raster van 3 of 4 grote afbeeldingenkaarten waar de speler uit kiest.", "2x2 of 1x3 kaartenraster", "Answer Options Grid"),
        ("CARD_CHOICE_DOLPHIN", "Keuzekaart: Dolfijn", "Afbeelding van een dolfijn (foutief antwoord bij 'krab').", "Witte kaart met grijze rand en illustratie", "Option 1 Selection"),
        ("CARD_CHOICE_CRAB", "Keuzekaart: Krab", "Afbeelding van een krab (correct antwoord op 'Waar is de krab?').", "Witte kaart met grijze rand en illustratie", "Option 2 (Correct)"),
        ("CARD_CHOICE_SHELLS", "Keuzekaart: Schelpen", "Afbeelding van schelpen en zeester (foutief antwoord).", "Witte kaart met grijze rand en illustratie", "Option 3 Selection"),
        ("FOOTER_QUIZ_PROGRESS", "Onderste Voortgangsbalk", "Voortgangsbalk voor 'Tempo' (0/10) en beloningsvoortgang (0/30 sterren).", "Gele afgeronde balk met groene progressbar", "Quiz Progress Footer")
    ]
    create_table(doc, std_headers, s6_data)

    # ---------------------------------------------------------
    # SECTION 9: SCHERM 7 - ZEG & ZET GAMEPLAY (SCR_ZEG_ZET_GAME)
    # ---------------------------------------------------------
    h1 = doc.add_heading(level=1)
    p_h1 = h1.paragraphs[0] if hasattr(h1, 'paragraphs') else h1
    r = p_h1.add_run("9. Scherm 7: Zeg & Zet Drag & Drop Gameplay (SCR_ZEG_ZET_GAME)")
    r.font.name = "Calibri"
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = TEAL
    p_h1.paragraph_format.space_before = Pt(14)
    p_h1.paragraph_format.space_after = Pt(6)

    add_screenshot_figure(doc, img_zeg_zet_game, "SCR_ZEG_ZET_GAME - Drag & Drop Opdracht Gameplay")
    s7_data = [
        ("BTN_NAV_BACK", "Terug Knop", "Verlaat de opdracht en keert terug naar het avontuurscherm.", "Witte cirkel knop met pijl-links", "Navigate Back"),
        ("DSP_STAR_COUNTER", "Sterrenteller", "Toont actuele sterrenstand ('0').", "Gele capsule pil met ster", "Score Display"),
        ("BTN_HINT_ASSIST", "Hint Knop", "Toont waar het object geplaatst moet worden.", "Gele knop met lamp-icoon", "Trigger Placement Hint"),
        ("BTN_ACTION_KLAAR", "Klaar Knop", "Valideert of de objecten op de juiste positie op het strand geplaatst zijn.", "Groene knop met vinkje-icoon en tekst 'Klaar'", "Validate Placement"),
        ("CARD_TASK_HEADER", "Opdracht Header Card", "Container met de te voltooien instructie en microfoon/toetsenbord knoppen.", "Witte afgeronde header met schaduw", "Task Instruction Header"),
        ("LBL_TASK_SENTENCE", "Opdrachtzin", "Instructie: 'Zet de boot in de zee.'", "Vetgedrukte instructietekst", "Placement Directive"),
        ("BTN_TASK_AUDIO", "Microfoon Knop", "Spreekt de opdracht in of luistert naar het antwoord.", "Groene cirkel knop met microfoon-icoon", "Voice Record Action"),
        ("BTN_TASK_KEYBOARD_TOGGLE", "Toetsenbord Knop", "Opent het typ-modal voor handmatige tekstinvoer (SCR_ZEG_ZET_KEYBOARD_OVERLAY).", "Blauwe cirkel knop met toetsenbord-icoon", "Open Keyboard Modal"),
        ("CANVAS_BEACH_OCEAN_SCENE", "Interactieve Strandscene", "Achtergrond met interactieve dropposities (zee, strand, palmeiland).", "Kleurrijke strand- en zee-art met dropzones", "Drag & Drop Canvas"),
        ("INFOBOX_VIDEO_ERROR_BANNER", "Fallback Melding", "Staat toe dat bij videostoringen de audio of tekst gelezen kan worden.", "Witte en gele afgeronde meldingskaart", "Fallback Notification"),
        ("TRAY_STICKER_PALETTE", "Sticker Keuzebalk", "Onderste carrousel met sleepbare objectstickers.", "Zachtgele afgeronde palletbalk", "Sticker Item Carousel"),
        ("STICKER_ITEMS", "Sleepbare Stickers", "Stickers: Dolfijn, Boot, Vuurtoren, Vliegtuig, Vlieger.", "Kleurrijke omrande afbeeldingen", "Draggable Sticker Objects"),
        ("BTN_TRAY_NEXT", "Volgende Stickers Knop", "Bladert naar de volgende set stickers in het palet.", "Witte cirkel met pijl-rechts", "Scroll Carousel")
    ]
    create_table(doc, std_headers, s7_data)

    # ---------------------------------------------------------
    # SECTION 10: SCHERM 8 - ZEG & ZET TYP DE ZIN OVERLAY (SCR_ZEG_ZET_KEYBOARD_OVERLAY)
    # ---------------------------------------------------------
    h1 = doc.add_heading(level=1)
    p_h1 = h1.paragraphs[0] if hasattr(h1, 'paragraphs') else h1
    r = p_h1.add_run("10. Scherm 8: Zeg & Zet Typ de Zin Modal Overlay (SCR_ZEG_ZET_KEYBOARD_OVERLAY)")
    r.font.name = "Calibri"
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = TEAL
    p_h1.paragraph_format.space_before = Pt(14)
    p_h1.paragraph_format.space_after = Pt(6)

    add_screenshot_figure(doc, img_zeg_zet_kbd, "SCR_ZEG_ZET_KEYBOARD_OVERLAY - Typ de Zin Modal Overlay")
    s8_data = [
        ("CARD_MODAL_KEYBOARD_INPUT", "Typ Modal Container", "Pop-up venster dat opent wanneer op de toetsenbordknop wordt getikt.", "Witte afgeronde kaart met blauwe rand en schaduw", "Modal Container"),
        ("BADGE_KEYBOARD_HEADER", "Toetsenbord Badge", "Blauwe cirkel met toetsenbord-icoon bovenaan de modal.", "Blauwe badge met toetsenbord-icoon", "Modal Header Icon"),
        ("TTL_KEYBOARD_MODAL", "Modal Titel", "Titel: 'TYP DE ZIN'.", "Donkerblauwe vetgedrukte titel", "Modal Header Text"),
        ("LBL_KEYBOARD_EXAMPLE", "Voorbeeldtekst", "Instructie: 'Bijvoorbeeld: Zet de boot in de zee.'", "Grijze instructietekst", "Input Guidance"),
        ("INPUT_SENTENCE_FIELD", "Tekstinvoerveld", "Invoerveld waarin de speler de zin kan typen ('Zet de boot in de zee.').", "Witte afgeronde rechthoek met blauwe rand", "Text Input Field"),
        ("BTN_KEYBOARD_SUBMIT", "Gebruik Zin Knop", "Bevestigt de getypte zin en verwerkt deze als antwoord.", "Brede blauwe knop met witte tekst 'Gebruik zin'", "Submit Input"),
        ("BTN_KEYBOARD_CLOSE", "Sluit Knop", "Sluit de typ-overlay zonder de voerwijziging op te slaan.", "Witte afgeronde knop met tekst 'Sluit'", "Dismiss Modal")
    ]
    create_table(doc, std_headers, s8_data)

    # ---------------------------------------------------------
    # SECTION 11: SCHERM 9 - ZEG & VLIEG ACTIEVE GAMEPLAY (SCR_ZEG_VLIEG_ACTIVE)
    # ---------------------------------------------------------
    h1 = doc.add_heading(level=1)
    p_h1 = h1.paragraphs[0] if hasattr(h1, 'paragraphs') else h1
    r = p_h1.add_run("11. Scherm 9: Zeg & Vlieg Actieve Vlieg Gameplay (SCR_ZEG_VLIEG_ACTIVE)")
    r.font.name = "Calibri"
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = TEAL
    p_h1.paragraph_format.space_before = Pt(14)
    p_h1.paragraph_format.space_after = Pt(6)

    add_screenshot_figure(doc, img_fly_active, "SCR_ZEG_VLIEG_ACTIVE - Actieve Bezem-Vlieg Gameplay Interface")
    s9_data = [
        ("BTN_NAV_HOME", "Home Knop", "Onderbreekt het vliegen en keert terug naar het menu.", "Witte vierkante knop met home-icoon", "Pause / Exit Game"),
        ("BAR_DISTANCE_PROGRESS", "Afstandsmeter", "Toont gevlogen meters ('Afstand 44m', geel gevulde balk 44/100).", "Blauwe capsule met gele progressbar", "Realtime Progress"),
        ("DSP_FLIGHT_SCORE_BADGE", "Score & Level Badge", "Toont trofee '44', sterren '0' en actueel level ('Level 1').", "Blauwe badge met trofee en ster-icoon", "Flight Stats Display"),
        ("CANVAS_FLIGHT_ARENA", "Vlieg Arena Canvas", "Scrollende 2D-wereld waarin het personage vliegt en obstakels ontwijkt.", "Dynamische strand- en luchtomgeving", "2D Physics World"),
        ("SPRITE_FLYING_HERO", "Speler Personage", "Jongen op de vliegende strandbezem (beweegt verticaal op spraak/knoppen).", "Geanimeerde spraakgestuurde avatar", "Player Avatar Sprite"),
        ("SPRITE_OBSTACLE_SEAGULL", "Obstakel: Meeuw", "Vliegende meeuw in de lucht die ontweken moet worden.", "Vogel sprite", "Air Obstacle Hazard"),
        ("SPRITE_OBSTACLE_SHARK", "Obstakel: Haai", "Springende haai uit het water.", "Haai sprite met ster-item", "Water Hazard"),
        ("SPRITE_COLLECTIBLE_ITEMS", "Verzamelbare Items", "Sterren en strandballen die extra punten opleveren.", "Glimmende sterren en strandbal sprites", "Bonus Collectible"),
        ("SLIDER_HEIGHT_CONTROL", "Hoogte-indicator Slider", "Verticale slider aan de rechterzijde die de actuele vlieghoogte toont.", "Transparante balk met rode positie-indicator", "Height Position Meter"),
        ("CARD_MIC_SPEECH_PROMPT", "Spraak Statusbalk", "Kaart onderaan: 'Noem wat je ziet' met microfoonstatus.", "Witte afgeronde container met mascotte", "Speech Recognition Dock"),
        ("BTN_FLY_UP", "Omhoog Knop", "Handmatige besturingsknop om te stijgen.", "Groene knop met pijl-omhoog 'Omhoog'", "Manual Altitude UP"),
        ("BTN_FLY_DOWN", "Omlaag Knop", "Handmatige besturingsknop om te dalen.", "Groene knop met pijl-omlaag 'Omlaag'", "Manual Altitude DOWN")
    ]
    create_table(doc, std_headers, s9_data)

    # ---------------------------------------------------------
    # SECTION 12: COMPONENT STATES & SFX MATRIX
    # ---------------------------------------------------------
    h1 = doc.add_heading(level=1)
    p_h1 = h1.paragraphs[0] if hasattr(h1, 'paragraphs') else h1
    r = p_h1.add_run("12. Component States & Geluidseffecten Matrix")
    r.font.name = "Calibri"
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = TEAL
    p_h1.paragraph_format.space_before = Pt(14)
    p_h1.paragraph_format.space_after = Pt(6)

    t12 = doc.add_table(rows=5, cols=4)
    t12.alignment = WD_TABLE_ALIGNMENT.CENTER
    t12.autofit = False
    t12.columns[0].width = Inches(1.5)
    t12.columns[1].width = Inches(1.8)
    t12.columns[2].width = Inches(1.7)
    t12.columns[3].width = Inches(1.5)
    
    t12_headers = ["State Name", "Visuele Transformatie", "Audio Effect (SFX)", "Haptische Feedback"]
    for i, h_text in enumerate(t12_headers):
        c = t12.rows[0].cells[i]
        set_cell_background(c, "0B8457")
        set_cell_margins(c, top=80, bottom=80, left=80, right=80)
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h_text)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9)

    state_data = [
        ("DEFAULT / IDLE", "Normale schaal 1.0x, standaard schaduw en helderheid.", "Geen geluid", "Geen trilling"),
        ("HOVER / FOCUS", "Subtiele vergroting (scale 1.05x), verhoogde schaduw, lichte glow.", "Zachte hoorbare tick / pop", "Geen trilling"),
        ("PRESSED / ACTIVE", "Gekrompen schaal (scale 0.95x), verlaagde schaduw (ingedrukt effect).", "Helder geluidsklik 'btn_click.mp3'", "Lichte haptische tik (5ms)"),
        ("DISABLED / LOCKED", "50% transparantie (opacity 0.5), slot-icoon overlay, geen hover-effect.", "Foutgeluid 'buzz_disabled.mp3'", "Korte dubbele trilling")
    ]

    for r_idx, row in enumerate(state_data, start=1):
        cells = t12.rows[r_idx].cells
        bg = "F4F7F6" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row):
            set_cell_background(cells[c_idx], bg)
            set_cell_margins(cells[c_idx], top=60, bottom=60, left=80, right=80)
            p = cells[c_idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(val)
            r.font.size = Pt(8.5)
            if c_idx == 0:
                r.font.bold = True
                r.font.color.rgb = TEAL

    set_table_borders(t12)
    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    add_callout_box(
        doc,
        [
            "✓ Alle 9 schermen en modals van Strand-bezem-escape zijn volledig gespecificeerd.",
            "✓ Frontend ontwikkelaars dienen de exacte Code Names (bijv. BTN_MANUAL_FLY_UP) te gebruiken in data-attributes en componentnamen.",
            "✓ QA-teams kunnen aan de hand van dit document alle schermen, knoppen, modals en gameplay-elementen testen."
        ],
        title="VOLLEDIGE ACCEPTATIECRITERIA & QA CHECKLIST",
        border_color="D98310",
        bg_color="FFF8E7"
    )

    output_path = "/Users/melkonian/git/Game-Wereld/docs/UX_Design_Specification_Magisch_Strand_Avontuur.docx"
    doc.save(output_path)
    print(f"Successfully generated full Word document with all 9 screens at {output_path}")

if __name__ == "__main__":
    build_docx()
